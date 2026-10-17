"""Trade-agnostic Excel and CSV exports from validated structured takeoff data."""
from __future__ import annotations

import csv
import io
from typing import Any, Dict, List, Union

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.schemas.takeoff import TakeoffExportRequest, TakeoffExtractionResponse


class ExportService:
    """Create exports from the data's own sections and terms; no trade templates."""

    def _data(self, data: Union[TakeoffExportRequest, TakeoffExtractionResponse, Dict[str, Any]]) -> Dict[str, Any]:
        if hasattr(data, "model_dump"):
            raw = data.model_dump(mode="json")
        else:
            raw = dict(data)
        items = raw.get("all_items") or raw.get("items") or []
        if not items:
            for section in (raw.get("sections") or {}).values():
                if isinstance(section, dict):
                    items.extend(section.get("items") or [])
        return {
            "metadata": raw.get("project_metadata") or {},
            "items": items,
            "sections": raw.get("sections") or {},
            "summary": raw.get("summary") or {},
            "terms": raw.get("proposal_terms") or {},
            "warnings": raw.get("warnings") or [],
            "filename": raw.get("filename") or "takeoff",
        }

    @staticmethod
    def _money(value: Any) -> Any:
        if value is None or value == "":
            return "Not stated — review source"
        try:
            return float(value)
        except (TypeError, ValueError):
            return str(value)

    def generate_excel(self, data: Union[TakeoffExportRequest, TakeoffExtractionResponse, Dict[str, Any]]) -> bytes:
        d = self._data(data)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Takeoff"
        commercial = wb.create_sheet("Commercial Terms")
        summary = wb.create_sheet("Document Summary")

        navy = PatternFill("solid", fgColor="1F4E78")
        blue = PatternFill("solid", fgColor="D9EAF7")
        light = PatternFill("solid", fgColor="F3F6F9")
        white_bold = Font(bold=True, color="FFFFFF", size=11)
        bold = Font(bold=True)
        thin = Side(style="thin", color="D5DDE5")
        border = Border(bottom=thin)

        meta = d["metadata"]
        title = f"{meta.get('trade') or 'Construction'} Takeoff Schedule"
        ws.append([title])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
        ws.cell(1, 1).fill = navy
        ws.cell(1, 1).font = Font(bold=True, color="FFFFFF", size=15)
        ws.cell(1, 1).alignment = Alignment(horizontal="center")
        for label, key in [("Project", "project_name"), ("Location", "location"), ("Contractor / Prepared by", "prepared_by"), ("Proposal date", "bid_date"), ("Trade", "trade")]:
            ws.append([label, meta.get(key) if meta.get(key) not in (None, "") else "Not stated — review source"])
        ws.append([])
        headers = ["Item #", "Section", "Description", "Model / Code", "Quantity", "Unit", "Source Page", "Review Status", "Source Evidence"]
        ws.append(headers)
        header_row = ws.max_row
        for cell in ws[header_row]:
            cell.fill = navy
            cell.font = white_bold
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for idx, item in enumerate(d["items"], 1):
            section = item.get("section") or "Uncategorized"
            evidence = item.get("source_evidence") or "Not available — review source"
            status = "Review Required" if item.get("review_required") else ("Approved" if item.get("approved") else "Extracted — verify")
            reasons = item.get("review_reasons") or []
            if reasons:
                status += ": " + "; ".join(str(x) for x in reasons)
            ws.append([
                item.get("item_number") or idx,
                section,
                item.get("original_description") or "Not stated — review source",
                item.get("model_number") or item.get("item_code") or "",
                item.get("quantity"),
                item.get("unit") or "",
                item.get("source_page"),
                status,
                evidence,
            ])
        for row in ws.iter_rows(min_row=header_row + 1):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            if str(row[7].value).startswith("Review Required"):
                row[7].font = Font(color="C00000", bold=True)
        widths = [12, 24, 48, 24, 13, 12, 13, 32, 60]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:I{ws.max_row}"

        commercial.append(["Commercial Terms and Proposal Conditions"])
        commercial.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
        commercial.cell(1, 1).fill = navy
        commercial.cell(1, 1).font = Font(bold=True, color="FFFFFF", size=14)
        commercial.append(["Category", "Description", "Quantity / Rate", "Amount / Value"])
        for cell in commercial[2]:
            cell.fill = navy
            cell.font = white_bold
        terms = d["terms"]
        pricing = [
            ("Pricing", "Subtotal", "", self._money(terms.get("subtotal"))),
            ("Pricing", "Tax rate (%)", "", self._money(terms.get("tax_rate"))),
            ("Pricing", "Tax amount", "", self._money(terms.get("tax_amount"))),
            ("Pricing", "Base proposal total", "", self._money(terms.get("base_proposal_total"))),
        ]
        for row in pricing:
            commercial.append(list(row))
        list_fields = [
            ("Option", "options"), ("Monitoring rate", "monitoring_rates"),
            ("Payment milestone", "payment_schedule"), ("Note / condition", "notes"),
            ("Customer responsibility", "customer_responsibilities"), ("Exclusion", "exclusions"),
            ("Installation condition", "installation_conditions"),
            ("Contractor responsibility", "contractor_responsibilities"),
            ("Warranty", "warranties"), ("General condition", "general_conditions"),
        ]
        for category, field in list_fields:
            for value in terms.get(field) or []:
                if isinstance(value, dict):
                    desc = value.get("description") or value.get("milestone") or ""
                    qty = value.get("quantity")
                    unit = value.get("unit")
                    amount = value.get("amount") if value.get("amount") is not None else value.get("amount_due")
                    if amount is None:
                        amount = value.get("amount")
                    commercial.append([category, desc, " ".join(str(x) for x in (qty, unit) if x not in (None, "")), self._money(amount) if amount is not None else ""])
                else:
                    commercial.append([category, str(value), "", ""])
        if len(commercial["A"]) <= 2:
            commercial.append(["Review", "No commercial terms were extracted; verify source PDF.", "", ""])
        for row in commercial.iter_rows(min_row=3):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for i, width in enumerate([26, 70, 24, 28], 1):
            commercial.column_dimensions[get_column_letter(i)].width = width
        commercial.freeze_panes = "A3"

        summary.append(["Document Summary"])
        summary.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        summary.cell(1, 1).fill = navy
        summary.cell(1, 1).font = Font(bold=True, color="FFFFFF", size=14)
        summary_rows = [
            ("Filename", d["filename"]),
            ("Project", meta.get("project_name") or "Not stated — review source"),
            ("Contractor / Prepared by", meta.get("prepared_by") or "Not stated — review source"),
            ("Trade", meta.get("trade") or "Not stated — review source"),
            ("Items extracted", len(d["items"])),
            ("Items with stated quantity", sum(i.get("quantity") is not None for i in d["items"])),
            ("Items missing quantity", sum(i.get("quantity") is None for i in d["items"])),
            ("Items requiring review", sum(bool(i.get("review_required")) for i in d["items"])),
        ]
        for row in summary_rows:
            summary.append(list(row))
        summary.append([])
        summary.append(["Document warnings"])
        for warning in d["warnings"]:
            summary.append([str(warning)])
        for row in summary.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.column == 1:
                    cell.font = bold
        summary.column_dimensions["A"].width = 32
        summary.column_dimensions["B"].width = 90

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def generate_csv(self, data: Union[TakeoffExportRequest, TakeoffExtractionResponse, Dict[str, Any]]) -> str:
        d = self._data(data)
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(["Record Type", "Item #", "Section / Category", "Description", "Model / Code", "Quantity", "Unit", "Amount / Value", "Source Page", "Review Status", "Source Evidence"])
        for idx, item in enumerate(d["items"], 1):
            status = "Review Required" if item.get("review_required") else ("Approved" if item.get("approved") else "Extracted — verify")
            if item.get("review_reasons"):
                status += ": " + "; ".join(str(x) for x in item["review_reasons"])
            writer.writerow(["Takeoff Item", item.get("item_number") or idx, item.get("section") or "Uncategorized", item.get("original_description"), item.get("model_number") or item.get("item_code"), item.get("quantity"), item.get("unit"), "", item.get("source_page"), status, item.get("source_evidence")])
        terms = d["terms"]
        for key, label in [("subtotal", "Subtotal"), ("tax_rate", "Tax rate (%)"), ("tax_amount", "Tax amount"), ("base_proposal_total", "Base proposal total")]:
            writer.writerow(["Pricing", "", "Pricing", label, "", "", "", self._money(terms.get(key)), "", "", ""])
        for category, field in [("Option", "options"), ("Monitoring rate", "monitoring_rates"), ("Payment milestone", "payment_schedule"), ("Note / condition", "notes"), ("Customer responsibility", "customer_responsibilities"), ("Exclusion", "exclusions"), ("Installation condition", "installation_conditions"), ("Contractor responsibility", "contractor_responsibilities"), ("Warranty", "warranties"), ("General condition", "general_conditions")]:
            for value in terms.get(field) or []:
                if isinstance(value, dict):
                    desc = value.get("description") or value.get("milestone") or ""
                    qty = value.get("quantity")
                    unit = value.get("unit")
                    amount = value.get("amount") if value.get("amount") is not None else value.get("amount_due")
                    writer.writerow([category, "", category, desc, "", qty, unit, amount, "", "", ""])
                else:
                    writer.writerow([category, "", category, str(value), "", "", "", "", "", "", ""])
        for warning in d["warnings"]:
            writer.writerow(["Review Warning", "", "Document warning", warning, "", "", "", "", "", "Review Required", ""])
        return output.getvalue()


export_service = ExportService()
