from io import BytesIO
import csv
import openpyxl

from app.schemas.takeoff import ProjectMetadata, TakeoffItem, TakeoffExportRequest, ProposalTerms
from app.services.export_service import export_service


def sample_payload():
    return TakeoffExportRequest(
        project_metadata=ProjectMetadata(project_name="Test Project", prepared_by="Example Contractor", trade="Electrical"),
        all_items=[
            TakeoffItem(
                item_number="1", section="lighting_schedule", original_description="LED fixture type A",
                model_number="ABC-123", quantity=12, unit="EA", source_page=2,
                source_evidence="Fixture schedule: 12 fixtures", review_required=False,
            ),
            TakeoffItem(
                item_number="2", section="conduit_and_wiring", original_description="Conduit routing shown",
                quantity=None, unit=None, source_page=None, source_evidence=None,
                review_required=True, review_reasons=["Quantity not stated"],
            ),
        ],
        proposal_terms=ProposalTerms(subtotal=None, tax_rate=None, tax_amount=None, base_proposal_total=None),
        filename="test-proposal.pdf",
    )


def test_excel_export_is_trade_agnostic_and_missing_total_is_not_zero():
    content = export_service.generate_excel(sample_payload())
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
    assert "Takeoff" in workbook.sheetnames
    assert "Commercial Terms" in workbook.sheetnames
    ws = workbook["Takeoff"]
    flattened = " ".join(str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None)
    assert "lighting_schedule" in flattened
    assert "conduit_and_wiring" in flattened
    assert "ABC-123" in flattened
    commercial = workbook["Commercial Terms"]
    commercial_text = " ".join(str(cell.value) for row in commercial.iter_rows() for cell in row if cell.value is not None)
    assert "Not stated — review source" in commercial_text
    assert "$0.00" not in commercial_text


def test_csv_export_preserves_dynamic_sections_and_missing_data():
    text = export_service.generate_csv(sample_payload())
    rows = list(csv.reader(text.splitlines()))
    flat = " ".join(cell for row in rows for cell in row)
    assert "lighting_schedule" in flat
    assert "conduit_and_wiring" in flat
    assert "ABC-123" in flat
    assert "Not stated — review source" in flat
